from pathlib import Path
from openpyxl import load_workbook

p = Path(__file__).resolve().parent / "data" / "LOYS_Aktien_Challenge.xlsx"
wb = load_workbook(p)
ws = wb["Teilnahmen"]
print("Teilnahmen:", max(ws.max_row - 1, 0))
print("Excel-Datei lesbar:", p)
wb.close()
