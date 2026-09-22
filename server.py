from http.server import ThreadingHTTPServer, SimpleHTTPRequestHandler
from pathlib import Path
from urllib.parse import urlparse
from openpyxl import load_workbook
from datetime import datetime
from threading import Lock
import shutil
import json
import re
import os
import sys
import unicodedata

ROOT = Path(__file__).resolve().parent
TEMPLATE_DATA_FILE = ROOT / "data" / "LOYS_Aktien_Challenge.xlsx"
DATA_DIR = Path(os.environ.get("DATA_DIR", str(ROOT / "data"))).resolve()
DATA_FILE = DATA_DIR / "LOYS_Aktien_Challenge.xlsx"
WEB_DIR = ROOT / "web"
LOCK = Lock()
HOST = os.environ.get("HOST", "127.0.0.1")
PORT = int(os.environ.get("PORT", "8000"))

DATA_DIR.mkdir(parents=True, exist_ok=True)
if DATA_FILE != TEMPLATE_DATA_FILE and not DATA_FILE.exists():
    shutil.copy2(TEMPLATE_DATA_FILE, DATA_FILE)

MAG7_TICKERS = {"AAPL", "MSFT", "NVDA", "GOOGL", "GOOG", "AMZN", "META", "TSLA"}


def normalize_name(value):
    value = unicodedata.normalize("NFKD", str(value or ""))
    value = "".join(ch for ch in value if not unicodedata.combining(ch))
    value = value.lower().replace("&", " and ")
    value = re.sub(r"[^a-z0-9]+", " ", value)
    return re.sub(r"\s+", " ", value).strip()


def normalize_ticker(value):
    value = str(value or "").strip().upper()
    value = re.sub(r"\s+", " ", value)
    return value


def ticker_root(value):
    value = normalize_ticker(value)
    return re.split(r"[\s\.:/\-]+", value, maxsplit=1)[0] if value else ""


def is_mag7_name(value):
    n = normalize_name(value)
    if not n:
        return False
    exact = {"apple", "microsoft", "nvidia", "alphabet", "google", "amazon", "meta", "facebook", "tesla"}
    if n in exact:
        return True
    prefixes = (
        "apple inc", "apple incorporated", "apple computer",
        "microsoft corp", "microsoft corporation",
        "nvidia corp", "nvidia corporation",
        "alphabet inc", "google llc",
        "amazon com", "amazon inc",
        "meta platforms", "facebook inc",
        "tesla inc", "tesla motors",
    )
    return any(n.startswith(p) for p in prefixes)


def is_mag7_pick(pick):
    return ticker_root(pick.get("ticker")) in MAG7_TICKERS or is_mag7_name(pick.get("name"))


def clean_pick(raw):
    raw = raw if isinstance(raw, dict) else {}
    name = str(raw.get("name", "")).strip()
    ticker = normalize_ticker(raw.get("ticker", ""))
    return {"name": name, "ticker": ticker}


def validate_pick(pick, label):
    name = pick["name"]
    ticker = pick["ticker"]
    if not name:
        return f"Bitte den Aktiennamen für {label} eingeben."
    if len(name) > 120:
        return f"Der Aktienname für {label} ist zu lang."
    if not ticker:
        return f"Bitte den Ticker für {label} eingeben."
    if len(ticker) > 24 or not re.match(r"^[A-Z0-9][A-Z0-9 .:/\-]{0,23}$", ticker):
        return f"Bitte einen gültigen Ticker für {label} eingeben."
    if is_mag7_pick(pick):
        return "Aktien der Magnificent Seven sind in dieser Challenge ausgeschlossen."
    return None


def read_json(handler):
    length = int(handler.headers.get("Content-Length", 0))
    raw = handler.rfile.read(length) if length else b"{}"
    return json.loads(raw.decode("utf-8"))


def send_json(handler, payload, status=200):
    body = json.dumps(payload, ensure_ascii=False).encode("utf-8")
    handler.send_response(status)
    handler.send_header("Content-Type", "application/json; charset=utf-8")
    handler.send_header("Content-Length", str(len(body)))
    handler.end_headers()
    handler.wfile.write(body)


def participation_rows(wb):
    ws = wb["Teilnahmen"]
    rows = []
    for r in ws.iter_rows(min_row=2, values_only=True):
        if not r or not r[0]:
            continue
        rows.append({"id": r[0], "email": r[3]})
    return rows


class Handler(SimpleHTTPRequestHandler):
    def translate_path(self, path):
        parsed = urlparse(path)
        rel = parsed.path.lstrip("/") or "index.html"
        target = (WEB_DIR / rel).resolve()
        if not str(target).startswith(str(WEB_DIR.resolve())):
            return str(WEB_DIR / "index.html")
        return str(target)

    def log_message(self, fmt, *args):
        sys.stdout.write("%s - %s\n" % (self.address_string(), fmt % args))

    def do_GET(self):
        path = urlparse(self.path).path
        if path == "/health":
            return send_json(self, {"ok": True})
        if path == "/api/rules":
            return send_json(self, {
                "mag7_excluded": True,
                "mag7": ["Apple", "Microsoft", "NVIDIA", "Alphabet", "Amazon", "Meta Platforms", "Tesla"],
                "free_text_entry": True,
            })
        return super().do_GET()

    def do_POST(self):
        path = urlparse(self.path).path
        if path != "/api/submit":
            return send_json(self, {"error": "Nicht gefunden"}, 404)

        try:
            payload = read_json(self)
        except Exception:
            return send_json(self, {"error": "Ungültige Anfrage"}, 400)

        name = str(payload.get("name", "")).strip()
        email = str(payload.get("email", "")).strip().lower()
        longs = [clean_pick(x) for x in payload.get("longs", [])]
        down = clean_pick(payload.get("down", {}))
        consent = bool(payload.get("consent"))

        if not name or len(name) > 80:
            return send_json(self, {"error": "Bitte Name oder Kürzel angeben."}, 400)
        if not email or not re.match(r"^[^@\s]+@[^@\s]+\.[^@\s]+$", email):
            return send_json(self, {"error": "Bitte gültige E-Mail-Adresse angeben."}, 400)
        if len(longs) != 3:
            return send_json(self, {"error": "Bitte genau drei Aktien für steigende Kurse eingeben."}, 400)

        for i, pick in enumerate(longs, start=1):
            error = validate_pick(pick, f"Favorit {i}")
            if error:
                return send_json(self, {"error": error}, 400)
        error = validate_pick(down, "die Aktie mit erwarteten fallenden Kursen")
        if error:
            return send_json(self, {"error": error}, 400)

        all_picks = longs + [down]
        ticker_keys = [normalize_ticker(p["ticker"]) for p in all_picks]
        name_keys = [normalize_name(p["name"]) for p in all_picks]
        if len(set(ticker_keys)) != 4:
            return send_json(self, {"error": "Bitte vier unterschiedliche Aktien eingeben. Ein Ticker wurde mehrfach verwendet."}, 400)
        if len(set(name_keys)) != 4:
            return send_json(self, {"error": "Bitte vier unterschiedliche Aktien eingeben. Ein Aktienname wurde mehrfach verwendet."}, 400)
        if not consent:
            return send_json(self, {"error": "Bitte Teilnahmebedingungen bestätigen."}, 400)

        with LOCK:
            wb = load_workbook(DATA_FILE)
            existing = participation_rows(wb)
            if any(str(p.get("email", "")).lower() == email for p in existing):
                wb.close()
                return send_json(self, {"error": "Diese E-Mail-Adresse hat bereits teilgenommen."}, 409)

            ws = wb["Teilnahmen"]
            next_no = max(ws.max_row, 1)
            pid = f"LC-{datetime.now().strftime('%Y%m%d')}-{next_no:04d}"
            now = datetime.now().strftime("%d.%m.%Y %H:%M:%S")

            ws.append([
                pid,
                now,
                name,
                email,
                longs[0]["name"],
                longs[0]["ticker"],
                longs[1]["name"],
                longs[1]["ticker"],
                longs[2]["name"],
                longs[2]["ticker"],
                down["name"],
                down["ticker"],
                "Ja",
                "Bestätigt",
            ])
            wb.save(DATA_FILE)
            wb.close()

        return send_json(self, {"ok": True, "id": pid})


if __name__ == "__main__":
    os.chdir(WEB_DIR)
    print(f"LOYS Aktien-Challenge läuft auf http://{HOST}:{PORT}")
    print(f"Excel-Datei: {DATA_FILE}")
    ThreadingHTTPServer((HOST, PORT), Handler).serve_forever()
